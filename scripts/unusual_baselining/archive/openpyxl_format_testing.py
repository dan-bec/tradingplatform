from openpyxl import Workbook
from openpyxl.styles import PatternFill, Border, Side, Font, NamedStyle
from openpyxl.utils import get_column_letter
from pathlib import Path

# Create a workbook and select the active worksheet
wb = Workbook()
ws = wb.active

# Example data (replace with your actual data)
ws['A1'] = "Header 1"
ws['B1'] = "Header 2"
ws['C1'] = "Header 3"
ws['A2'] = 10
ws['B2'] = 20
ws['C2'] = 30
ws['A3'] = 40
ws['B3'] = 50
ws['C3'] = 60

# 1. Set background colors for table headers
header_fill = PatternFill(start_color='DDDDDD', end_color='DDDDDD', fill_type='solid')  # Light gray
for cell in ws['A1':'C1'][0]:  # First row
    cell.fill = header_fill
    cell.font = Font(bold=True)  # Bold headers

# 2. Add borders to the table (A1:C3)
thin_border = Border(left=Side(style='thin'), right=Side(style='thin'), 
                     top=Side(style='thin'), bottom=Side(style='thin'))
for row in ws['A1':'C3']:
    for cell in row:
        cell.border = thin_border

# 3. Adjust column widths automatically based on content
for col in ws.columns:
    max_length = 0
    column = col[0].column_letter
    for cell in col:
        try:
            if len(str(cell.value)) > max_length:
                max_length = len(str(cell.value))
        except:
            pass
    adjusted_width = (max_length + 2)  # Add padding
    ws.column_dimensions[column].width = adjusted_width

# 4. Add a notes section with formatting
ws['A5'] = "Notes:"
ws['A5'].font = Font(bold=True)
ws['A6'] = "This is a note about the data."
notes_fill = PatternFill(start_color='F0F0F0', end_color='F0F0F0', fill_type='solid')  # Very light gray
for row in ws['A5':'A6']:
    for cell in row:
        cell.fill = notes_fill

# Save the workbook
wb.save(Path(__file__).parent / "formatted_output.xlsx")