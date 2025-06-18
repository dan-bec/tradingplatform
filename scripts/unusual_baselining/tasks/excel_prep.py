import sys
from pathlib import Path

# Determine the project root dynamically
TASK_SCRIPT_DIR = Path(__file__).parent
REPO_ROOT = TASK_SCRIPT_DIR.parents[2]  

# Insert the project root into sys.path if not already present
if str(REPO_ROOT) not in sys.path:
    sys.path.insert(0, str(REPO_ROOT))
    
import duckdb
import pandas as pd
from openpyxl import Workbook
from openpyxl.utils.dataframe import dataframe_to_rows
from openpyxl.utils import get_column_letter
from openpyxl.styles import PatternFill, Border, Side, Font, NamedStyle
from datetime import datetime
import scripts.unusual_baselining.config as config
import argparse

# Parse command-line arguments
parser = argparse.ArgumentParser(description="Generate Excel files for each security.")
parser.add_argument("--output-dir", type=str, default=config.COMPILED_OUTPUT_DIR / "security_summary", help="Directory to save Excel files")
args = parser.parse_args()

# Ensure output directory exists
output_dir = Path(args.output_dir)
output_dir.mkdir(parents=True, exist_ok=True)

# Connect to DuckDB
con = duckdb.connect(config.FULL_DB_PATH)

# Get unique securities
compiled_schema = config.COMPILED
itm_threshold_100 = config.ITM_THRESHOLD_100
table_name = f"{compiled_schema}.all_options_trades_above_baseline_{itm_threshold_100}"
securities = con.execute(f"SELECT DISTINCT security FROM {table_name} WHERE security = 'ABBV'").fetchall()
securities = [sec[0] for sec in securities]

# Define percentage style for Excel
percentage_style = NamedStyle(name="percentage", number_format="0.00%")

# Define the win/loss column name (adjust this based on your findings)
WL_COLUMN = 'itm'  # Replace with the actual column name, e.g., 'win_loss'

for security in securities:
    # Query data for the security
    query = f"SELECT * FROM {table_name} WHERE security = ?"
    df = con.execute(query, [security]).fetchdf()
    
    if df.empty:
        print(f"No data for {security}, skipping.")
        continue
    
    # Print columns for debugging
    print(f"Columns in DataFrame for {security}: {df.columns.tolist()}")
    
    # Check if the win/loss column exists
    if WL_COLUMN not in df.columns:
        print(f"Error: Column '{WL_COLUMN}' not found in DataFrame for {security}. Skipping.")
        continue
    
    # Compute aggregates for Summary tab
    total_trades = len(df)
    calls = df[df["option_type"] == "C"]
    puts = df[df["option_type"] == "P"]
    calls_count = len(calls)
    puts_count = len(puts)
    calls_wins = (calls[WL_COLUMN] == 1).sum()
    puts_wins = (puts[WL_COLUMN] == 1).sum()
    calls_losses = (calls[WL_COLUMN] == 0).sum()
    puts_losses = (puts[WL_COLUMN] == 0).sum()
    calls_win_pct = calls_wins / calls_count if calls_count > 0 else 0
    puts_win_pct = puts_wins / puts_count if puts_count > 0 else 0
    
    # Types of Trades
    trade_types = df.groupby("option_condition_name").agg(
        count=("option_ticker", "count"),
        wins=(WL_COLUMN, lambda x: (x == 1).sum()),
        losses=(WL_COLUMN, lambda x: (x == 0).sum())
    ).reset_index()
    trade_types["win_pct"] = trade_types["wins"] / trade_types["count"]
    
    # Expiration Month
    df["expiration_month"] = pd.to_datetime(df["expiration"]).dt.to_period("M")
    expiration_months = df.groupby("expiration_month").agg(
        count=("option_ticker", "count"),
        wins=(WL_COLUMN, lambda x: (x == 1).sum()),
        losses=(WL_COLUMN, lambda x: (x == 0).sum())
    ).reset_index()
    expiration_months["win_pct"] = expiration_months["wins"] / expiration_months["count"]
    expiration_months["expiration_month"] = expiration_months["expiration_month"].astype(str)
    df = df.drop(columns=["expiration_month"])

    # Create Excel file
    wb = Workbook()
    ws_summary = wb.active
    ws_summary.title = "Summary"
    ws_raw = wb.create_sheet("Raw Data")
    
    # Write Summary tab
    ws_summary["A1"] = f"{security} Analysis"
    ws_summary["A2"] = f"Completed: {datetime.now().strftime('%m/%d/%y')}"
    ws_summary["A4"] = "Summary Data"
    summary_data = pd.DataFrame({
        "": ["Total # of Trades", "Calls", "Puts"],
        "Count": [total_trades, calls_count, puts_count],
        "Win": ["", calls_wins, puts_wins],
        "Loss": ["", calls_losses, puts_losses],
        "Win %": ["", calls_win_pct, puts_win_pct]
    })
    for r, row in enumerate(dataframe_to_rows(summary_data, index=False, header=True), start=5):
        for c, val in enumerate(row, start=1):
            cell = ws_summary.cell(row=r, column=c, value=val)
            if c == 5 and r > 5:  # Apply percentage style to Win % column
                cell.style = percentage_style
    
    # Write Types of Trades
    ws_summary["A10"] = "Types of Trades"
    for r, row in enumerate(dataframe_to_rows(trade_types, index=False, header=True), start=11):
        for c, val in enumerate(row, start=1):
            cell = ws_summary.cell(row=r, column=c, value=val)
            if c == 5 and r > 11:  # Apply percentage style to win_pct column
                cell.style = percentage_style
    
    # Write Expiration Month
    ws_summary["A20"] = "Expiration Month"
    for r, row in enumerate(dataframe_to_rows(expiration_months, index=False, header=True), start=21):
        for c, val in enumerate(row, start=1):
            cell = ws_summary.cell(row=r, column=c, value=val)
            if c == 5 and r > 21:  # Apply percentage style to win_pct column
                cell.style = percentage_style
    
    # Write Raw Data tab
    for r, row in enumerate(dataframe_to_rows(df, index=False, header=True), start=1):
        for c, val in enumerate(row, start=1):
            ws_raw.cell(row=r, column=c, value=val)
    
    # Save the Excel file
    file_name = output_dir / f"{security}_analysis.xlsx"
    wb.save(file_name)
    print(f"Saved {file_name}")

# Close DuckDB connection
con.close()